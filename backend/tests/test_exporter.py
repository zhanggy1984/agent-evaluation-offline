"""6.1 报告导出纯逻辑单测（宿主直跑，零 DB 依赖）。

render_* 是纯函数（payload dict → bytes），本测试覆盖：
- PDF/Excel 魔数与结构（可被下载端识别）
- 水印确实进入产物（PDF 字节随水印变化；Excel 可读回页脚）
- 维度分 / 错误 / 数值格式化（报告数据正确性）
- token 派生：sha256 hex 长度 == export_token.token_hash CHAR(64)

一次性消费 / 过期 / user_id 绑定 / viewer 403 等 HTTP 行为由容器 e2e 覆盖（6.1c）。
"""
import io
import unittest
from hashlib import sha256

from openpyxl import load_workbook

from app.exporter.render import (
    _dims, _err, _fmt, _num, _sanitize_cell, render_pdf, render_xlsx,
)

RUN = {
    "run_id": 41, "agent_name": "customer-service", "version": "v2.1.0",
    "trigger_type": "manual", "status": "completed",
    "started_at": "2026-08-17 10:00:00", "finished_at": "2026-08-17 10:05:00",
    "agent_score": 79.17, "total_case": 2, "pass_case": 1, "fail_case": 1,
    "error_case": 0, "na_case": 0,
    "ttft_p50": 120.5, "ttft_p95": 200.1, "e2e_p50": 1500.3, "e2e_p95": 2600.0,
    "total_tokens": 2832, "total_cost": 0.004867, "models": ["deepseek-chat"],
}

RESULTS = [
    {
        "case_name": "查询订单状态", "interface_name": "/api/chat", "pass_fail": "pass",
        "score_total": 95.5,
        "score_per_dimension": [
            {"code": "completeness", "value": 100.0, "na": False},
            {"code": "token_cost", "value": 1.2, "na": False},
        ],
        "error_type": None, "error_detail": None,
        "model": "deepseek-chat", "total_tokens": 1500, "total_cost": 0.0025,
    },
    {
        "case_name": "退款流程", "interface_name": "/api/chat", "pass_fail": "fail",
        "score_total": 55.0,
        "score_per_dimension": [
            {"code": "completeness", "value": 55.0, "na": False},
            {"code": "factuality", "na": True, "na_reason": "judge 超时"},
        ],
        "error_type": "assert_fail", "error_detail": "断言 expected≠actual",
        "model": "deepseek-chat", "total_tokens": 1332, "total_cost": 0.002367,
    },
]

PAYLOAD = {"run": RUN, "results": RESULTS}


def _open_pdf(data: bytes) -> None:
    assert data[:4] == b"%PDF", "PDF 魔数缺失"


class TestRenderPdf(unittest.TestCase):
    def test_magic_bytes(self):
        _open_pdf(render_pdf(PAYLOAD, "admin @ 2026-08-18 10:00:00"))

    def test_empty_results_renders(self):
        _open_pdf(render_pdf({"run": RUN, "results": []}, "w"))

    def test_watermark_affects_output(self):
        """水印不同 → 字节不同（证明水印进入产物）。"""
        a = render_pdf(PAYLOAD, "alice @ 2026-08-18 10:00:00")
        b = render_pdf(PAYLOAD, "bob @ 2026-08-18 10:00:00")
        self.assertNotEqual(a, b)


class TestRenderXlsx(unittest.TestCase):
    def test_sheets_and_rows(self):
        wb = load_workbook(io.BytesIO(render_xlsx(PAYLOAD, "admin @ 2026-08-18 10:00:00")))
        self.assertEqual(wb.sheetnames, ["汇总", "用例明细", "成本明细"])
        # 汇总含 agent 名；用例明细含两条；成本明细含合计行
        self.assertEqual(wb["汇总"].cell(2, 2).value, "customer-service")
        self.assertEqual(wb["用例明细"].max_row, 3)
        self.assertEqual(wb["成本明细"].cell(4, 1).value, "合计")

    def test_footer_watermark(self):
        """Excel 页脚水印可读回（openpyxl 直读）。"""
        wb = load_workbook(io.BytesIO(render_xlsx(PAYLOAD, "admin @ 2026-08-18")))
        self.assertEqual(wb["汇总"].oddFooter.center.text, "admin @ 2026-08-18")

    def test_empty_results_renders(self):
        wb = load_workbook(io.BytesIO(render_xlsx({"run": RUN, "results": []}, "w")))
        self.assertEqual(wb["用例明细"].max_row, 1)


class TestFormatters(unittest.TestCase):
    def test_fmt(self):
        self.assertEqual(_fmt(79.1667, 2), 79.17)
        self.assertEqual(_fmt(0.004867, 6), 0.004867)
        self.assertIsNone(_fmt(None))

    def test_num(self):
        self.assertEqual(_num(None), "N/A")
        self.assertEqual(_num(79.17, 2), "79.17")
        self.assertEqual(_num(2832, 0), "2832")

    def test_dims(self):
        s = _dims(RESULTS[0]["score_per_dimension"])
        self.assertIn("完成度: 100.00", s)
        self.assertIn("Token 成本: 1.20", s)
        s2 = _dims([{"code": "factuality", "na": True, "na_reason": "judge 超时"}])
        self.assertIn("事实性: N/A（judge 超时）", s2)

    def test_err(self):
        self.assertEqual(_err(RESULTS[1]), "assert_fail: 断言 expected≠actual")
        self.assertEqual(_err({"error_type": None, "error_detail": None}), "")


class TestTokenHash(unittest.TestCase):
    def test_sha256_hex_fits_char64(self):
        """export_token.token_hash CHAR(64) == sha256 hex（64 字符）。"""
        h = sha256(b"token").hexdigest()
        self.assertEqual(len(h), 64)


class TestInjectionSanitize(unittest.TestCase):
    """7.6 A6 导出注入净化：公式前缀 / HTML 转义。"""

    def test_sanitize_cell_prefixes_formula(self):
        for payload in ("=1+1", "+cmd", "-2+3", "@SUM(A1)", "\t=x", "\r\n=1"):
            self.assertTrue(_sanitize_cell(payload).startswith("'"), payload)
        self.assertEqual(_sanitize_cell("normal text"), "normal text")
        self.assertEqual(_sanitize_cell("1+1"), "1+1")   # 非公式开头的数字串不误伤
        self.assertEqual(_sanitize_cell(123), 123)        # 数值原样（openpyxl 写数字）

    def test_xlsx_formula_mitigated(self):
        """case_name / agent_name / error_detail 以 = @ 开头 → 导出后单元格以 ' 开头（公式不执行）。"""
        p = {
            "run": {**RUN, "agent_name": "=cmd()"},
            "results": [{**RESULTS[0], "case_name": "=1+1", "error_detail": "@Eval('x')"}],
        }
        wb = load_workbook(io.BytesIO(render_xlsx(p, "w")))
        self.assertTrue(str(wb["汇总"].cell(2, 2).value).startswith("'"))       # agent_name
        self.assertTrue(str(wb["用例明细"].cell(2, 2).value).startswith("'"))   # case_name
        self.assertTrue(str(wb["用例明细"].cell(2, 7).value).startswith("'"))   # error_detail 列

    def test_pdf_escapes_html(self):
        """error_detail 含 <&> → PDF 正常渲染（Paragraph 已转义，不抛异常）。"""
        p = {
            "run": RUN,
            "results": [{**RESULTS[1], "error_detail": "x<y & z>w"}],
        }
        data = render_pdf(p, "w")
        self.assertTrue(data[:4] == b"%PDF")


if __name__ == "__main__":
    unittest.main()
